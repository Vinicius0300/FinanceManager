"""Exportação do histórico de operações para Excel (.xlsx), com os mesmos filtros da tela.

Decisões do dono (Parte 4): planilha Excel em vez de CSV; valores gravados como **número**
(o Excel mostra o separador decimal do Windows) e datas como **data** no formato
``AAAA-MM-DD``. O arquivo é escrito no caminho escolhido no diálogo nativo do Tauri:
primeiro num temporário ao lado e depois trocado, para não deixar um .xlsx pela metade.

Os números vão para a célula já arredondados (dinheiro 2 casas; quantidade, preço e câmbio
até 8), convertidos de ``Decimal`` só na escrita: o Excel guarda tudo como ponto flutuante
de 15 dígitos, suficiente para valores com essas casas.
"""

import os
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.orm import Session

from app.asset_types.catalogo import REGISTRO_PADRAO
from app.asset_types.registry import RegistroTiposAtivo
from app.domain.erros import DetalheCampo, OperacaoNaoPermitidaError, ValidacaoError
from app.domain.historico import SituacaoPosicao
from app.domain.numeros import CASAS_PRECO, arredondar, arredondar_moeda
from app.services.historico import (
    CampoOrdem,
    Direcao,
    FiltroHistorico,
    ItemHistorico,
    consultar,
    resumir,
)

FORMATO_DATA = "yyyy-mm-dd"
FORMATO_MOEDA = "#,##0.00"
FORMATO_FINO = "#,##0.########"

NOMES_OPERACAO = {
    "compra": "Compra",
    "venda": "Venda",
    "aporte": "Aporte",
    "resgate": "Resgate",
    "provento": "Provento",
}
NOMES_SITUACAO = {
    SituacaoPosicao.ABERTA: "Ainda na carteira",
    SituacaoPosicao.ENCERRADA: "Posição encerrada",
}

# (título, largura, formato)
COLUNAS: tuple[tuple[str, int, str | None], ...] = (
    ("Data", 12, FORMATO_DATA),
    ("Operação", 11, None),
    ("Ativo", 14, None),
    ("Nome", 28, None),
    ("Tipo de ativo", 16, None),
    ("Instituição", 20, None),
    ("Moeda", 7, None),
    ("Quantidade", 14, FORMATO_FINO),
    ("Preço unitário", 15, FORMATO_FINO),
    ("Taxas", 11, FORMATO_MOEDA),
    ("Imposto retido", 14, FORMATO_MOEDA),
    ("Câmbio", 11, FORMATO_FINO),
    ("Valor total (moeda)", 18, FORMATO_MOEDA),
    ("Valor total (BRL)", 18, FORMATO_MOEDA),
    ("Resultado (moeda)", 18, FORMATO_MOEDA),
    ("Resultado (BRL)", 18, FORMATO_MOEDA),
    ("Situação do ativo", 18, None),
    ("Observação", 30, None),
)


@dataclass(frozen=True)
class ArquivoExportado:
    caminho: Path
    linhas: int


def _numero(valor: Decimal | None, casas: int | None = None) -> float | None:
    if valor is None:
        return None
    arredondado = arredondar_moeda(valor) if casas is None else arredondar(valor, casas)
    return float(arredondado)


def _linha(item: ItemHistorico, registro: RegistroTiposAtivo) -> list[object]:
    t, op = item.transacao, item.operacao
    tipo_ativo = registro.obter(item.ativo.tipo).nome if item.ativo.tipo in registro else ""
    resultado = op.resultado
    return [
        t.data,
        NOMES_OPERACAO[t.tipo.value],
        item.ativo.simbolo or "",
        item.ativo.nome,
        tipo_ativo,
        item.instituicao.nome,
        t.moeda,
        _numero(t.quantidade, CASAS_PRECO),
        _numero(t.preco_unitario, CASAS_PRECO),
        _numero(t.taxas),
        _numero(t.imposto_retido),
        _numero(t.cambio, CASAS_PRECO),
        _numero(op.valor_total),
        _numero(op.valor_total_brl),
        None if resultado is None else _numero(resultado.resultado),
        None if resultado is None else _numero(resultado.resultado_brl),
        NOMES_SITUACAO[item.situacao],
        t.observacao or "",
    ]


def _formatar(planilha: Worksheet, primeira_linha: int) -> None:
    for indice, (_, largura, formato) in enumerate(COLUNAS, start=1):
        letra = get_column_letter(indice)
        planilha.column_dimensions[letra].width = largura
        if formato is None:
            continue
        for (celula,) in planilha.iter_rows(min_row=primeira_linha, min_col=indice, max_col=indice):
            celula.number_format = formato


def _aba_resumo(planilha: Worksheet, sessao: Session, filtro: FiltroHistorico, hoje: date) -> None:
    resumo = resumir(sessao, filtro, hoje)
    totais = resumo.totais
    planilha.append(["Resumo do filtro", None])
    planilha["A1"].font = Font(bold=True)
    linhas: list[tuple[str, object]] = [
        ("Período (início)", resumo.inicio),
        ("Período (fim)", resumo.fim),
        ("Operações", totais.operacoes),
        ("Total comprado (compras + aportes, BRL)", _numero(totais.total_comprado)),
        ("Total vendido (vendas + resgates, BRL)", _numero(totais.total_vendido)),
        ("Lucro realizado (BRL)", _numero(totais.lucro_realizado)),
        ("Proventos líquidos (BRL)", _numero(totais.proventos_liquidos)),
        ("IR retido nos proventos (BRL)", _numero(totais.impostos_retidos)),
        ("Taxas pagas (BRL)", _numero(totais.taxas)),
        ("Operações com taxas", totais.operacoes_com_taxas),
    ]
    for rotulo, valor in linhas:
        planilha.append([rotulo, valor])
        celula = planilha.cell(row=planilha.max_row, column=2)
        if isinstance(valor, date):
            celula.number_format = FORMATO_DATA
        elif isinstance(valor, float):
            celula.number_format = FORMATO_MOEDA
    planilha.column_dimensions["A"].width = 42
    planilha.column_dimensions["B"].width = 18


def validar_caminho(caminho: str) -> Path:
    """Caminho absoluto terminado em .xlsx, numa pasta que existe."""
    destino = Path(caminho)
    if not destino.is_absolute():
        raise ValidacaoError(
            "Escolha onde salvar a planilha.",
            [DetalheCampo("caminho", "Informe um caminho completo (ex.: C:\\...\\arquivo.xlsx).")],
        )
    if destino.suffix.lower() != ".xlsx":
        raise ValidacaoError(
            "O arquivo precisa ter a extensão .xlsx.",
            [DetalheCampo("caminho", "Use a extensão .xlsx.")],
        )
    if not destino.parent.is_dir():
        raise ValidacaoError(
            "A pasta escolhida não existe.",
            [DetalheCampo("caminho", "Escolha uma pasta existente.")],
        )
    return destino


def exportar_historico(
    sessao: Session,
    filtro: FiltroHistorico,
    caminho: str,
    hoje: date,
    campo: CampoOrdem = CampoOrdem.DATA,
    direcao: Direcao = Direcao.DESC,
    registro: RegistroTiposAtivo = REGISTRO_PADRAO,
) -> ArquivoExportado:
    """Grava a planilha (aba "Operações" + aba "Resumo") e devolve o caminho e as linhas."""
    destino = validar_caminho(caminho)
    itens = consultar(sessao, filtro, campo, direcao, registro).itens

    livro = Workbook()
    operacoes = livro.active
    if operacoes is None:  # pragma: no cover - Workbook novo sempre tem uma aba
        operacoes = livro.create_sheet()
    operacoes.title = "Operações"
    operacoes.append([titulo for titulo, _, _ in COLUNAS])
    for celula in operacoes[1]:
        celula.font = Font(bold=True)
    operacoes.freeze_panes = "A2"
    for item in itens:
        operacoes.append(_linha(item, registro))
    _formatar(operacoes, primeira_linha=2)
    _aba_resumo(livro.create_sheet("Resumo"), sessao, filtro, hoje)

    temporario = destino.with_name(f".{destino.stem}.tmp.xlsx")
    try:
        livro.save(temporario)
        os.replace(temporario, destino)
    except PermissionError as erro:
        temporario.unlink(missing_ok=True)
        raise OperacaoNaoPermitidaError(
            "Não foi possível salvar a planilha. Se ela estiver aberta no Excel, feche-a e "
            "tente de novo."
        ) from erro
    except OSError as erro:
        temporario.unlink(missing_ok=True)
        raise OperacaoNaoPermitidaError(
            f"Não foi possível salvar a planilha: {erro.strerror or 'erro de gravação'}."
        ) from erro
    return ArquivoExportado(destino, len(itens))
