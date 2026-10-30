"""API do histórico de operações: filtros, ordem, resumo e exportação (Parte 4).

Mesmo cenário de ``test_historico.py``, montado pela API (relógio fixo em 29/10/2026):
PETR4 100 na XP a R$ 30 e 50 na RICO a R$ 30; as 50 da RICO vendidas a R$ 34 com R$ 5 de
taxa (+R$ 195); dividendo de R$ 85 líquidos com R$ 15 de IR na XP.
"""

from pathlib import Path
from typing import Any

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from tests.test_api import api, app_api, criar, yahoo  # noqa: F401 - fixtures reutilizadas

Json = dict[str, Any]


def _cenario(api: TestClient) -> Json:  # noqa: F811
    xp = criar(api, "/api/instituicoes", {"nome": "XP", "tipo": "corretora"})["id"]
    rico = criar(api, "/api/instituicoes", {"nome": "RICO", "tipo": "corretora"})["id"]
    for inst, valor in ((xp, "3000"), (rico, "1500")):
        criar(
            api,
            "/api/movimentos",
            {"instituicao_id": inst, "tipo": "deposito", "data": "2026-01-02", "valor": valor},
        )
    petr = criar(api, "/api/ativos", {"tipo": "acao_br", "nome": "Petrobras", "simbolo": "PETR4"})
    base = {"ativo_id": petr["id"], "tipo": "compra", "preco_unitario": "30"}
    c1 = criar(
        api,
        "/api/transacoes",
        {**base, "instituicao_id": xp, "data": "2026-01-10", "quantidade": "100"},
    )
    c2 = criar(
        api,
        "/api/transacoes",
        {**base, "instituicao_id": rico, "data": "2026-01-12", "quantidade": "50"},
    )
    venda = {
        "ativo_id": petr["id"],
        "instituicao_id": rico,
        "tipo": "venda",
        "data": "2026-03-05",
        "quantidade": "50",
        "preco_unitario": "34",
        "taxas": "5",
    }
    v1 = criar(api, "/api/transacoes", venda)
    provento = {
        "ativo_id": petr["id"],
        "instituicao_id": xp,
        "tipo": "provento",
        "data": "2026-04-15",
        "valor": "85",
        "imposto_retido": "15",
        "subtipo_provento": "dividendo",
    }
    p1 = criar(api, "/api/transacoes", provento)
    return {"xp": xp, "rico": rico, "petr": petr["id"], "c1": c1, "c2": c2, "v1": v1, "p1": p1}


def test_saida_traz_valor_total_resultado_e_situacao(api: TestClient) -> None:  # noqa: F811
    c = _cenario(api)
    venda = c["v1"]
    assert venda["resultado"] == "195.00"
    assert venda["resultado_brl"] == "195.00"
    assert venda["valor_total"] == "1695.00"
    assert venda["valor_total_brl"] == "1695.00"
    assert venda["taxas_brl"] == "5.00"
    assert venda["situacao"] == "encerrada"  # posição da RICO zerada
    assert c["c1"]["resultado"] is None
    assert c["c1"]["valor_total"] == "3000.00"
    assert c["p1"]["imposto_retido_brl"] == "15.00"
    assert c["p1"]["valor_total"] == "85.00"
    detalhe = api.get(f"/api/transacoes/{venda['id']}").json()
    assert detalhe["resultado_brl"] == "195.00"


def test_listagem_com_filtros_ordem_e_paginacao(api: TestClient) -> None:  # noqa: F811
    c = _cenario(api)

    def ids(**params: Any) -> list[int]:
        resposta = api.get("/api/transacoes", params=params)
        assert resposta.status_code == 200, resposta.text
        return [t["id"] for t in resposta.json()["itens"]]

    todos = ids()
    assert todos == [c["p1"]["id"], c["v1"]["id"], c["c2"]["id"], c["c1"]["id"]]
    assert ids(tipos=["venda", "provento"]) == [c["p1"]["id"], c["v1"]["id"]]
    assert ids(tipo="compra", tipos=["compra", "venda"]) == [c["c2"]["id"], c["c1"]["id"]]
    assert ids(busca="petrobras", moeda="brl", tipo_ativo="acao_br") == todos
    assert ids(tipo_ativo="cripto") == []
    assert ids(apenas_com_custos="true") == [c["p1"]["id"], c["v1"]["id"]]
    assert ids(apenas_prejuizo="true") == []
    # Situação: sem instituição, PETR4 está aberta (100 na XP); na RICO, encerrada.
    assert ids(situacao="aberta") == todos
    assert ids(situacao="encerrada", instituicao_id=c["rico"]) == [c["v1"]["id"], c["c2"]["id"]]
    assert ids(situacao="aberta", instituicao_id=c["rico"]) == []
    assert ids(inicio="2026-03-01", fim="2026-03-31") == [c["v1"]["id"]]
    assert ids(ordem="data", direcao="asc") == todos[::-1]
    assert ids(ordem="resultado")[0] == c["v1"]["id"]
    assert ids(ordem="valor_total", direcao="desc")[0] == c["c1"]["id"]

    pagina = api.get("/api/transacoes", params={"limite": 3, "deslocamento": 3}).json()
    assert pagina["total"] == 4
    assert [t["id"] for t in pagina["itens"]] == [c["c1"]["id"]]


def test_filtros_invalidos_em_pt_br(api: TestClient) -> None:  # noqa: F811
    resposta = api.get("/api/transacoes", params={"inicio": "2026-05-01", "fim": "2026-04-01"})
    assert resposta.status_code == 422
    corpo = resposta.json()
    assert corpo["mensagem"] == "A data inicial não pode ser depois da final."
    assert corpo["detalhes"][0]["campo"] == "inicio"
    assert api.get("/api/transacoes", params={"tipos": ["doacao"]}).status_code == 422
    assert api.get("/api/transacoes", params={"situacao": "talvez"}).status_code == 422
    assert api.get("/api/transacoes", params={"ordem": "cor"}).status_code == 422
    assert api.get("/api/transacoes/resumo", params={"moeda": "REAL"}).status_code == 422


def test_resumo(api: TestClient) -> None:  # noqa: F811
    c = _cenario(api)
    resposta = api.get("/api/transacoes/resumo", params={"inicio": "2025-11-01"})
    assert resposta.status_code == 200, resposta.text
    resumo = resposta.json()
    assert resumo["totais"] == {
        "operacoes": 4,
        "total_comprado": "4500.00",
        "total_vendido": "1695.00",
        "lucro_realizado": "195.00",
        "proventos_liquidos": "85.00",
        "impostos_retidos": "15.00",
        "taxas": "5.00",
        "operacoes_com_taxas": 1,
    }
    assert resumo["total_sem_filtros"] == 4
    assert resumo["inicio"] == "2025-11-01"
    assert resumo["fim"] == "2026-10-29"
    assert len(resumo["serie"]) == 12
    janeiro = resumo["serie"][2]
    assert janeiro == {
        "mes": "2026-01-01",
        "compras": "4500.00",
        "vendas": "0.00",
        "proventos": "0.00",
    }
    (grupo,) = resumo["por_ativo"]
    assert grupo["ativo_id"] == c["petr"]
    assert grupo["situacao"] == "aberta"
    situacoes = {i["instituicao_id"]: i["situacao"] for i in grupo["instituicoes"]}
    assert situacoes == {c["xp"]: "aberta", c["rico"]: "encerrada"}

    so_rico = api.get("/api/transacoes/resumo", params={"instituicao_id": c["rico"]}).json()
    assert so_rico["totais"]["lucro_realizado"] == "195.00"
    assert so_rico["por_ativo"][0]["situacao"] == "encerrada"
    assert so_rico["total_sem_filtros"] == 4


def test_exportacao(api: TestClient, tmp_path: Path) -> None:  # noqa: F811
    _cenario(api)
    destino = tmp_path / "historico.xlsx"
    resposta = api.post(
        "/api/transacoes/exportar",
        params={"tipos": ["venda"]},
        json={"caminho": str(destino)},
    )
    assert resposta.status_code == 200, resposta.text
    assert resposta.json() == {"caminho": str(destino), "linhas": 1}
    linhas = list(load_workbook(destino)["Operações"].iter_rows(min_row=2, values_only=True))
    assert len(linhas) == 1
    assert linhas[0][1] == "Venda"

    invalido = api.post("/api/transacoes/exportar", json={"caminho": "planilha.xlsx"})
    assert invalido.status_code == 422
    assert invalido.json()["detalhes"][0]["campo"] == "caminho"


def test_rotas_do_historico_exigem_token(app_api: FastAPI) -> None:  # noqa: F811
    with TestClient(app_api) as sem_token:
        assert sem_token.get("/api/transacoes/resumo").status_code == 401
        assert sem_token.post("/api/transacoes/exportar", json={"caminho": "x"}).status_code == 401
