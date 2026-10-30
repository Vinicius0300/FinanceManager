"""Histórico de operações (Parte 4): resultado por venda/resgate, situação, filtros, resumo.

Todos os números conferidos à mão (contas nos comentários). Cenário base dos testes de
service, decidido com o dono para explicar as regras:

- PETR4: 100 compradas na XP a R$ 30 (custo R$ 3.000) e 50 na RICO a R$ 30 (R$ 1.500);
  as 50 da RICO vendidas a R$ 34 com R$ 5 de taxa → 1.700 − 5 − 1.500 = **+R$ 195**.
- AAPL (USD) na AVENUE: 10 a US$ 100 + US$ 1 de taxa, câmbio 5,00 (custo US$ 1.001 =
  R$ 5.005); vendidas a US$ 110 − US$ 1, câmbio 5,50 → US$ 1.099 = R$ 6.044,50.
  Resultado: **+US$ 98** e **+R$ 1.039,50** (inclui a alta do dólar).
- Dividendo de PETR4 na XP: R$ 85 líquidos com R$ 15 de IR retido.
"""

from collections.abc import Callable
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.domain.enums import TipoEvento, TipoTransacao
from app.domain.erros import OperacaoNaoPermitidaError, ValidacaoError
from app.domain.historico import (
    OperacaoHistorico,
    ResultadoOperacao,
    SituacaoPosicao,
    agrupar_por_ativo,
    analisar_cotas,
    analisar_montante,
    meses_entre,
    serie_mensal,
    situacao_do_ativo,
    totalizar,
    valor_total,
)
from app.domain.montante import Aporte, OperacaoMontante, Resgate
from app.domain.posicao import (
    Compra,
    EventoCorporativo,
    OperacaoCotas,
    Provento,
    Venda,
    aplicar_operacoes,
)
from app.repositories import modelos as m
from app.repositories.transacoes import RepositorioEventos, RepositorioTransacoes
from app.services import exportacao, historico
from app.services.historico import CampoOrdem, Direcao, FiltroHistorico
from app.services.posicoes import estado_cotas
from app.services.transacoes import ResultadoTransacao
from tests import fabricas as f
from tests.fabricas import D

A = SituacaoPosicao.ABERTA
E = SituacaoPosicao.ENCERRADA


# ------------------------------------------------------------------------- domínio puro


def test_resultado_de_cada_venda_pelo_preco_medio() -> None:
    # 100 a R$ 10 + 100 a R$ 12 → custo 2.200, PM 11.
    # Venda 1: 50 a R$ 15 com R$ 5 de taxa → 750 − 5 − 50 × 11 = +195.
    # Venda 2: 150 a R$ 9 → 1.350 − 150 × 11 = −300 (zera a posição).
    ops: list[OperacaoCotas] = [
        Compra(date(2026, 1, 5), D(100), D(10), ordem=1),
        Compra(date(2026, 2, 5), D(100), D(12), ordem=2),
        Venda(date(2026, 3, 5), D(50), D(15), taxas=D(5), ordem=3),
        Venda(date(2026, 4, 5), D(150), D(9), ordem=4),
    ]
    analise = analisar_cotas(ops)
    assert analise.resultados == {
        3: ResultadoOperacao(D(195), D(195)),
        4: ResultadoOperacao(D(-300), D(-300)),
    }
    assert analise.situacao is E
    # A soma dos resultados é o lucro realizado acumulado da posição (−105).
    soma = sum(r.resultado for r in analise.resultados.values())
    assert soma == aplicar_operacoes(ops).moeda.lucro_realizado == D(-105)


def test_resultado_em_moeda_estrangeira_inclui_o_cambio() -> None:
    # Custo US$ 1.001 (R$ 5.005); venda US$ 1.099 a 5,50 = R$ 6.044,50.
    ops: list[OperacaoCotas] = [
        Compra(date(2026, 1, 5), D(10), D(100), taxas=D(1), cambio=D(5), ordem=1),
        Venda(date(2026, 6, 5), D(10), D(110), taxas=D(1), cambio=D("5.5"), ordem=2),
    ]
    resultado = analisar_cotas(ops).resultados[2]
    assert resultado.resultado == D(98)
    assert resultado.resultado_brl == D("1039.5")


def test_resultado_respeita_desdobramento_e_ignora_provento() -> None:
    # 10 a R$ 20 (custo 200); desdobramento 1:2 → 20 ações, PM 10;
    # provento não mexe no custo; venda de 10 a R$ 12 → 120 − 100 = +20; sobram 10.
    ops: list[OperacaoCotas] = [
        Compra(date(2026, 1, 5), D(10), D(20), ordem=1),
        EventoCorporativo(date(2026, 2, 1), TipoEvento.DESDOBRAMENTO, 1, 2, ordem=1),
        Provento(date(2026, 2, 10), D(30), ordem=2),
        Venda(date(2026, 3, 5), D(10), D(12), ordem=3),
    ]
    analise = analisar_cotas(ops)
    assert analise.resultados == {3: ResultadoOperacao(D(20), D(20))}
    assert analise.situacao is A


def test_resultado_do_resgate_proporcional() -> None:
    # Fator fixo de 10%: aporte 1.000 vale 1.100. Resgate de 550 = metade do saldo:
    # sai metade do custo (500) → +50. Resgate do resto (550) → +50 e zera.
    def fator(inicio: date, fim: date) -> Decimal:
        return D("1.1") if fim > inicio else D(1)

    ops: list[OperacaoMontante] = [
        Aporte(date(2026, 1, 5), D(1000), ordem=1),
        Resgate(date(2026, 2, 5), D(550), ordem=2),
        Resgate(date(2026, 3, 5), D(550), ordem=3),
    ]
    analise = analisar_montante(ops, fator)
    assert analise.resultados[2] == ResultadoOperacao(D(50), D(50))
    assert analise.resultados[3] == ResultadoOperacao(D(50), D(50))
    assert analise.situacao is E
    assert analisar_montante(ops[:2], fator).situacao is A


def test_situacao_do_ativo_e_aberta_se_alguma_posicao_estiver() -> None:
    assert situacao_do_ativo([E, A]) is A
    assert situacao_do_ativo([E, E]) is E
    assert situacao_do_ativo([]) is E


def test_valor_total_e_o_movimento_de_caixa() -> None:
    # Prévia aprovada: 200 × 38,12 + 2 = 7.626; 150 × 52,40 − 3,50 = 7.856,50.
    compra = valor_total(
        TipoTransacao.COMPRA, quantidade=D(200), preco_unitario=D("38.12"), taxas=D(2), valor=None
    )
    venda = valor_total(
        TipoTransacao.VENDA,
        quantidade=D(150),
        preco_unitario=D("52.40"),
        taxas=D("3.5"),
        valor=None,
    )
    provento = valor_total(
        TipoTransacao.PROVENTO, quantidade=None, preco_unitario=None, taxas=D(0), valor=D(85)
    )
    assert (compra, venda, provento) == (D("7626.00"), D("7856.50"), D(85))


def _op(
    ident: int,
    dia: date,
    tipo: TipoTransacao,
    total: str,
    *,
    ativo: int = 1,
    inst: int = 1,
    taxas: str = "0",
    imposto: str = "0",
    cambio: str = "1",
    resultado: str | None = None,
) -> OperacaoHistorico:
    return OperacaoHistorico(
        id=ident,
        data=dia,
        tipo=tipo,
        ativo_id=ativo,
        instituicao_id=inst,
        cambio=D(cambio),
        valor_total=D(total),
        taxas=D(taxas),
        imposto_retido=D(imposto),
        resultado=None if resultado is None else ResultadoOperacao(D(resultado), D(resultado)),
    )


OPS = [
    _op(1, date(2026, 1, 10), TipoTransacao.COMPRA, "1000", taxas="2"),
    _op(2, date(2026, 1, 20), TipoTransacao.APORTE, "500", ativo=2, inst=2),
    # Venda em USD: 200 × 5 = R$ 1.000; taxa US$ 1 = R$ 5; resultado já vem em BRL.
    _op(3, date(2026, 3, 5), TipoTransacao.VENDA, "200", taxas="1", cambio="5", resultado="-40"),
    _op(4, date(2026, 3, 9), TipoTransacao.PROVENTO, "85", imposto="15"),
    _op(5, date(2026, 3, 9), TipoTransacao.RESGATE, "300", ativo=2, inst=2, resultado="20"),
]


def test_totais_do_filtro() -> None:
    totais = totalizar(OPS)
    assert totais.operacoes == 5
    assert totais.total_comprado == D(1500)  # 1.000 + 500
    assert totais.total_vendido == D(1300)  # 1.000 (USD × 5) + 300
    assert totais.lucro_realizado == D(-20)  # −40 + 20
    assert totais.proventos_liquidos == D(85)
    assert totais.impostos_retidos == D(15)  # opção A: fora das taxas
    assert totais.taxas == D(7)  # 2 + 1 × 5
    assert totais.operacoes_com_taxas == 2  # o provento com IR não conta
    assert totalizar([]).operacoes == 0


def test_operacao_marca_prejuizo_e_custos() -> None:
    assert OPS[2].prejuizo
    assert not OPS[4].prejuizo
    assert not OPS[0].prejuizo
    assert OPS[0].tem_custos
    assert OPS[3].tem_custos
    assert not OPS[1].tem_custos
    assert OPS[2].valor_total_brl == D(1000)
    assert OPS[2].taxas_brl == D(5)


def test_meses_entre_cruza_o_ano_e_aceita_dias_quaisquer() -> None:
    assert meses_entre(date(2025, 11, 30), date(2026, 2, 3)) == [
        date(2025, 11, 1),
        date(2025, 12, 1),
        date(2026, 1, 1),
        date(2026, 2, 1),
    ]
    assert meses_entre(date(2026, 3, 1), date(2026, 2, 1)) == []


def test_serie_mensal_inclui_meses_vazios() -> None:
    serie = serie_mensal(OPS, date(2026, 1, 1), date(2026, 4, 30))
    assert [p.mes.month for p in serie] == [1, 2, 3, 4]
    janeiro, fevereiro, marco, abril = serie
    assert (janeiro.compras, janeiro.vendas, janeiro.proventos) == (D(1500), D(0), D(0))
    assert (fevereiro.compras, fevereiro.vendas) == (D(0), D(0))
    assert (marco.compras, marco.vendas, marco.proventos) == (D(0), D(1300), D(85))
    assert abril.compras == D(0)
    # Operações fora do período não criam barras.
    assert serie_mensal(OPS, date(2026, 2, 1), date(2026, 2, 28))[0].vendas == D(0)


def test_agrupar_por_ativo_com_instituicoes() -> None:
    ops = [
        _op(1, date(2026, 1, 10), TipoTransacao.COMPRA, "3000", ativo=7, inst=1),
        _op(2, date(2026, 1, 11), TipoTransacao.COMPRA, "1500", ativo=7, inst=2),
        _op(3, date(2026, 2, 1), TipoTransacao.VENDA, "1695", ativo=7, inst=2, resultado="195"),
        _op(4, date(2026, 2, 2), TipoTransacao.COMPRA, "10", ativo=8, inst=1),
    ]
    situacoes = {(7, 1): A, (7, 2): E, (8, 1): A}
    grupos = agrupar_por_ativo(ops, lambda a, i: situacoes[(a, i)], lambda a: A)
    assert [g.ativo_id for g in grupos] == [7, 8]
    petr = grupos[0]
    assert petr.totais.total_comprado == D(4500)
    assert petr.totais.lucro_realizado == D(195)
    assert [(i.instituicao_id, i.situacao) for i in petr.instituicoes] == [(1, A), (2, E)]
    assert petr.instituicoes[1].totais.total_vendido == D(1695)
    assert petr.instituicoes[0].totais.lucro_realizado == D(0)


# --------------------------------------------------------------------------- services

Cenario = dict[str, m.Instituicao | m.Ativo | m.Transacao]


def _gravada(resultado: ResultadoTransacao) -> m.Transacao:
    assert resultado.transacao is not None
    return resultado.transacao


@pytest.fixture
def cenario(sessao: Session) -> Cenario:
    xp = f.instituicao(sessao, "XP")
    rico = f.instituicao(sessao, "RICO")
    avenue = f.instituicao(sessao, "Avenue")
    petr = f.acao_br(sessao, "PETR4")
    aapl = f.acao_us(sessao, "AAPL")
    cdb = f.cdb(sessao, indexador="pre", taxa="10", aplicacao=date(2026, 1, 2))
    c1 = _gravada(f.compra(sessao, petr, xp, date(2026, 1, 10), "100", "30"))
    c2 = _gravada(f.compra(sessao, petr, rico, date(2026, 1, 12), "50", "30"))
    v1 = _gravada(f.venda(sessao, petr, rico, date(2026, 3, 5), "50", "34", taxas="5"))
    p1 = _gravada(f.provento(sessao, petr, xp, date(2026, 4, 15), "85", imposto="15"))
    c3 = _gravada(f.compra(sessao, aapl, avenue, date(2026, 2, 2), "10", "100", "1", cambio="5"))
    v2 = _gravada(f.venda(sessao, aapl, avenue, date(2026, 9, 1), "10", "110", "1", cambio="5.5"))
    a1 = _gravada(f.montante(sessao, cdb, xp, TipoTransacao.APORTE, date(2026, 1, 2), "1000"))
    sessao.flush()
    itens: Cenario = {"xp": xp, "rico": rico, "avenue": avenue, "petr": petr, "aapl": aapl}
    itens |= {"cdb": cdb, "c1": c1, "c2": c2, "v1": v1, "p1": p1, "c3": c3, "v2": v2, "a1": a1}
    return itens


def _id(cenario: Cenario, chave: str) -> int:
    return cenario[chave].id


def _ids(consulta: historico.ConsultaHistorico) -> list[int]:
    return [i.transacao.id for i in consulta.itens]


def test_resultados_nas_vendas_e_soma_igual_ao_lucro_do_ativo(
    sessao: Session, cenario: Cenario
) -> None:
    itens = {i.transacao.id: i for i in historico.consultar(sessao, FiltroHistorico()).itens}
    venda_petr = itens[_id(cenario, "v1")].operacao
    venda_aapl = itens[_id(cenario, "v2")].operacao
    assert venda_petr.resultado == ResultadoOperacao(D(195), D(195))
    assert venda_aapl.resultado == ResultadoOperacao(D(98), D("1039.5"))
    assert venda_aapl.valor_total == D(1099)
    assert venda_aapl.valor_total_brl == D("6044.5")
    assert itens[_id(cenario, "c1")].operacao.resultado is None
    assert itens[_id(cenario, "p1")].operacao.resultado is None

    # Conferência pedida no prompt: soma dos resultados de cada ativo = lucro realizado.
    repo = RepositorioTransacoes(sessao)
    for chave in ("petr", "aapl"):
        ativo_id = _id(cenario, chave)
        eventos = RepositorioEventos(sessao).do_ativo(ativo_id)
        lucro_moeda = lucro_brl = D(0)
        for inst in {t.instituicao_id for t in repo.do_ativo(ativo_id)}:
            estado = estado_cotas(repo.da_posicao(ativo_id, inst), eventos)
            lucro_moeda += estado.moeda.lucro_realizado
            lucro_brl += estado.brl.lucro_realizado
        resultados = [
            i.operacao.resultado for i in itens.values() if i.transacao.ativo_id == ativo_id
        ]
        assert sum((r.resultado for r in resultados if r), D(0)) == lucro_moeda
        assert sum((r.resultado_brl for r in resultados if r), D(0)) == lucro_brl


def test_situacao_por_ativo_sem_filtro_de_instituicao(sessao: Session, cenario: Cenario) -> None:
    # Sem filtro: PETR4 ainda tem 100 na XP → todas as operações dela são "aberta",
    # inclusive a venda na RICO; AAPL foi toda vendida → "encerrada".
    consulta = historico.consultar(sessao, FiltroHistorico())
    situacoes = {i.transacao.id: i.situacao for i in consulta.itens}
    assert situacoes[_id(cenario, "v1")] is A
    assert situacoes[_id(cenario, "v2")] is E
    abertas = historico.consultar(sessao, FiltroHistorico(situacao=A))
    assert _id(cenario, "v1") in _ids(abertas)
    assert _id(cenario, "v2") not in _ids(abertas)
    encerradas = historico.consultar(sessao, FiltroHistorico(situacao=E))
    assert sorted(_ids(encerradas)) == sorted([_id(cenario, "c3"), _id(cenario, "v2")])


def test_situacao_por_posicao_com_filtro_de_instituicao(sessao: Session, cenario: Cenario) -> None:
    # Com filtro na RICO: a posição de PETR4 lá foi zerada → "encerrada".
    rico = _id(cenario, "rico")
    consulta = historico.consultar(sessao, FiltroHistorico(instituicao_id=rico))
    assert {i.situacao for i in consulta.itens} == {E}
    assert consulta.situacao_ativo(_id(cenario, "petr")) is E
    filtro_aberta = FiltroHistorico(instituicao_id=rico, situacao=A)
    assert historico.consultar(sessao, filtro_aberta).itens == []
    xp = FiltroHistorico(instituicao_id=_id(cenario, "xp"), situacao=A)
    assert len(historico.consultar(sessao, xp).itens) == 3  # compra, provento e aporte do CDB


@pytest.mark.parametrize(
    ("filtro", "esperados"),
    [
        (FiltroHistorico(tipo_ativo="acao_us"), ["c3", "v2"]),
        (FiltroHistorico(moeda="USD"), ["c3", "v2"]),
        (FiltroHistorico(busca="petr"), ["c1", "c2", "v1", "p1"]),
        (FiltroHistorico(busca="TESTE AAPL"), ["c3", "v2"]),  # nome "Teste AAPL"
        (FiltroHistorico(busca="banco x"), ["a1"]),  # nome do CDB
        (FiltroHistorico(busca="%"), []),  # curinga do LIKE é texto literal
        (FiltroHistorico(tipos=frozenset({TipoTransacao.VENDA})), ["v1", "v2"]),
        (
            FiltroHistorico(tipos=frozenset({TipoTransacao.COMPRA, TipoTransacao.APORTE})),
            ["c1", "c2", "c3", "a1"],
        ),
        (FiltroHistorico(inicio=date(2026, 3, 1), fim=date(2026, 4, 30)), ["v1", "p1"]),
        (FiltroHistorico(apenas_com_custos=True), ["v1", "p1", "c3", "v2"]),
        (FiltroHistorico(apenas_prejuizo=True), []),
    ],
)
def test_filtros(
    sessao: Session, cenario: Cenario, filtro: FiltroHistorico, esperados: list[str]
) -> None:
    obtidos = sorted(_ids(historico.consultar(sessao, filtro)))
    assert obtidos == sorted(_id(cenario, chave) for chave in esperados)


def test_filtro_so_vendas_com_prejuizo(sessao: Session, cenario: Cenario) -> None:
    # Nova venda de 50 PETR4 da XP a R$ 20: 1.000 − 50 × 30 = −500.
    petr, xp = cenario["petr"], cenario["xp"]
    assert isinstance(petr, m.Ativo)
    assert isinstance(xp, m.Instituicao)
    perda = f.venda(sessao, petr, xp, date(2026, 5, 2), "50", "20")
    assert perda.transacao is not None
    consulta = historico.consultar(sessao, FiltroHistorico(apenas_prejuizo=True))
    assert _ids(consulta) == [perda.transacao.id]
    assert consulta.itens[0].operacao.resultado == ResultadoOperacao(D(-500), D(-500))


def test_tipos_efetivos_combina_tipo_e_tipos() -> None:
    venda, compra = TipoTransacao.VENDA, TipoTransacao.COMPRA
    assert historico.tipos_efetivos(None, None) is None
    assert historico.tipos_efetivos(None, []) is None
    assert historico.tipos_efetivos(venda, None) == frozenset({venda})
    assert historico.tipos_efetivos(venda, [venda, compra]) == frozenset({venda})
    assert historico.tipos_efetivos(venda, [compra]) == frozenset()


def test_ordenacao(sessao: Session, cenario: Cenario) -> None:
    def ordem(campo: CampoOrdem, direcao: Direcao) -> list[int]:
        return _ids(historico.consultar(sessao, FiltroHistorico(), campo, direcao))

    datas = ordem(CampoOrdem.DATA, Direcao.DESC)
    assert datas[0] == _id(cenario, "v2")
    assert datas[-1] == _id(cenario, "a1")
    assert ordem(CampoOrdem.DATA, Direcao.ASC) == datas[::-1]
    # Resultado: só vendas têm valor; as demais vão para o fim nas duas direções.
    assert ordem(CampoOrdem.RESULTADO, Direcao.DESC)[:2] == [_id(cenario, "v2"), _id(cenario, "v1")]
    assert ordem(CampoOrdem.RESULTADO, Direcao.ASC)[:2] == [_id(cenario, "v1"), _id(cenario, "v2")]
    assert (
        ordem(CampoOrdem.RESULTADO, Direcao.ASC)[2:]
        == ordem(CampoOrdem.RESULTADO, Direcao.DESC)[2:]
    )
    # Valor total em BRL: a venda da AAPL (R$ 6.044,50) é a maior.
    assert ordem(CampoOrdem.VALOR_TOTAL, Direcao.DESC)[0] == _id(cenario, "v2")
    # Ativo por símbolo; o CDB (sem símbolo) usa o nome "CDB Banco X".
    por_ativo = ordem(CampoOrdem.ATIVO, Direcao.ASC)
    assert por_ativo[:2] == [_id(cenario, "v2"), _id(cenario, "c3")]  # AAPL, mais recente antes
    assert por_ativo[2] == _id(cenario, "a1")
    assert ordem(CampoOrdem.QUANTIDADE, Direcao.ASC)[-2:] == [
        _id(cenario, "p1"),
        _id(cenario, "a1"),
    ]


def test_resumo_totais_serie_e_por_ativo(sessao: Session, cenario: Cenario) -> None:
    hoje = date(2026, 10, 29)
    resumo = historico.resumir(sessao, FiltroHistorico(inicio=date(2025, 11, 1)), hoje)
    t = resumo.totais
    assert t.operacoes == 7
    assert resumo.total_sem_filtros == 7
    # Comprado: 3.000 + 1.500 + 1.001 × 5 + 1.000 = 10.505.
    assert t.total_comprado == D(10505)
    # Vendido: 1.695 + 1.099 × 5,5 = 7.739,50.
    assert t.total_vendido == D("7739.5")
    assert t.lucro_realizado == D("1234.5")  # 195 + 1.039,50
    assert (t.proventos_liquidos, t.impostos_retidos) == (D(85), D(15))
    assert t.taxas == D("15.5")  # 5 + 1 × 5 + 1 × 5,5
    assert t.operacoes_com_taxas == 3
    assert len(resumo.serie) == 12  # nov/2025 a out/2026
    assert resumo.serie[2].compras == D(5500)  # jan: 3.000 + 1.500 + 1.000
    assert resumo.serie[10].vendas == D("6044.5")  # set
    grupos = {g.ativo_id: g for g in resumo.por_ativo}
    petr = grupos[_id(cenario, "petr")]
    assert petr.situacao is A
    assert [(i.instituicao_id, i.situacao) for i in petr.instituicoes] == [
        (_id(cenario, "xp"), A),
        (_id(cenario, "rico"), E),
    ]
    assert petr.totais.lucro_realizado == D(195)
    assert grupos[_id(cenario, "aapl")].situacao is E


def test_resumo_sem_inicio_comeca_na_operacao_mais_antiga(
    sessao: Session, cenario: Cenario
) -> None:
    resumo = historico.resumir(sessao, FiltroHistorico(), date(2026, 10, 29))
    assert resumo.inicio == date(2026, 1, 2)
    assert resumo.serie[0].mes == date(2026, 1, 1)
    assert len(resumo.serie) == 10
    vazio = historico.resumir(sessao, FiltroHistorico(busca="nada"), date(2026, 10, 29))
    assert vazio.serie == []
    assert vazio.totais.operacoes == 0
    assert vazio.total_sem_filtros == 7


def test_resumo_com_filtro_de_instituicao(sessao: Session, cenario: Cenario) -> None:
    filtro = FiltroHistorico(instituicao_id=_id(cenario, "rico"))
    resumo = historico.resumir(sessao, filtro, date(2026, 10, 29))
    assert resumo.totais.lucro_realizado == D(195)
    (grupo,) = resumo.por_ativo
    assert grupo.situacao is E  # situação da posição na RICO, não do ativo


def test_operacao_da_transacao(sessao: Session, cenario: Cenario) -> None:
    venda = cenario["v1"]
    assert isinstance(venda, m.Transacao)
    operacao, situacao = historico.operacao_da_transacao(sessao, venda)
    assert operacao.resultado == ResultadoOperacao(D(195), D(195))
    assert situacao is E  # posição da própria instituição


def test_posicao_que_nao_fecha_nao_derruba_a_consulta(
    sessao: Session, cenario: Cenario, monkeypatch: pytest.MonkeyPatch
) -> None:
    from app.domain.erros import PosicaoInsuficienteError

    def falha(*_: object) -> None:
        raise PosicaoInsuficienteError("posição inválida")

    monkeypatch.setattr(historico, "analisar_cotas", falha)
    itens = historico.consultar(sessao, FiltroHistorico(tipo_ativo="acao_br")).itens
    assert len(itens) == 4
    assert all(i.operacao.resultado is None and i.situacao is A for i in itens)


# ---------------------------------------------------------------------------- exportação


def test_exporta_excel_com_numeros_e_datas(
    sessao: Session, cenario: Cenario, tmp_path: Path
) -> None:
    destino = tmp_path / "historico.xlsx"
    arquivo = exportacao.exportar_historico(
        sessao, FiltroHistorico(), str(destino), date(2026, 10, 29)
    )
    assert arquivo.linhas == 7
    assert arquivo.caminho == destino
    livro = load_workbook(destino)
    assert livro.sheetnames == ["Operações", "Resumo"]
    planilha = livro["Operações"]
    cabecalho = [c.value for c in planilha[1]]
    assert cabecalho[0] == "Data"
    assert "Resultado (BRL)" in cabecalho
    primeira = {cabecalho[i]: c for i, c in enumerate(planilha[2])}  # mais recente: venda AAPL
    assert primeira["Operação"].value == "Venda"
    assert primeira["Ativo"].value == "AAPL"
    assert primeira["Data"].value == datetime(2026, 9, 1)
    assert primeira["Data"].number_format == "yyyy-mm-dd"
    assert primeira["Valor total (BRL)"].value == pytest.approx(6044.5)
    assert primeira["Resultado (BRL)"].value == pytest.approx(1039.5)
    assert primeira["Resultado (moeda)"].value == pytest.approx(98)
    assert primeira["Câmbio"].value == pytest.approx(5.5)
    assert primeira["Situação do ativo"].value == "Posição encerrada"
    assert primeira["Tipo de ativo"].value  # nome do tipo vindo do registry
    resumo = {linha[0].value: linha[1].value for linha in livro["Resumo"].iter_rows(min_row=2)}
    assert resumo["Lucro realizado (BRL)"] == pytest.approx(1234.5)
    assert resumo["IR retido nos proventos (BRL)"] == pytest.approx(15)


def test_exportacao_respeita_filtros_e_ordem(
    sessao: Session, cenario: Cenario, tmp_path: Path
) -> None:
    destino = tmp_path / "vendas.xlsx"
    filtro = FiltroHistorico(tipos=frozenset({TipoTransacao.VENDA}))
    arquivo = exportacao.exportar_historico(
        sessao, filtro, str(destino), date(2026, 10, 29), CampoOrdem.DATA, Direcao.ASC
    )
    assert arquivo.linhas == 2
    linhas = list(load_workbook(destino)["Operações"].iter_rows(min_row=2, values_only=True))
    assert [linha[2] for linha in linhas] == ["PETR4", "AAPL"]


@pytest.mark.parametrize(
    ("caminho", "mensagem"),
    [
        ("relativo.xlsx", "Escolha onde salvar"),
        ("{tmp}/planilha.csv", "extensão .xlsx"),
        ("{tmp}/nao-existe/planilha.xlsx", "pasta escolhida não existe"),
    ],
)
def test_caminho_invalido(sessao: Session, tmp_path: Path, caminho: str, mensagem: str) -> None:
    with pytest.raises(ValidacaoError, match=mensagem):
        exportacao.exportar_historico(
            sessao, FiltroHistorico(), caminho.format(tmp=tmp_path), date(2026, 10, 29)
        )


@pytest.mark.parametrize(
    ("erro", "mensagem"),
    [
        (PermissionError(13, "negado"), "feche-a"),
        (OSError(28, "Disco cheio"), "Disco cheio"),
    ],
)
def test_falha_ao_gravar_vira_mensagem_e_limpa_o_temporario(
    sessao: Session,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    erro: OSError,
    mensagem: str,
) -> None:
    def recusar(*_: object) -> None:
        raise erro

    trocar: Callable[..., None] = recusar
    monkeypatch.setattr("app.services.exportacao.os.replace", trocar)
    with pytest.raises(OperacaoNaoPermitidaError, match=mensagem):
        exportacao.exportar_historico(
            sessao, FiltroHistorico(), str(tmp_path / "x.xlsx"), date(2026, 10, 29)
        )
    assert not any(p.name.endswith(".tmp.xlsx") for p in tmp_path.iterdir())
    assert not (tmp_path / "x.xlsx").exists()
